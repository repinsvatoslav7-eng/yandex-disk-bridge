import base64
import io
import os
import posixpath
import time
import traceback
from datetime import datetime, timezone
from typing import Any

import requests
from openpyxl import load_workbook
from mcp.server.fastmcp import FastMCP

YANDEX_TOKEN = os.environ["YANDEX_DISK_TOKEN"]
PORT = int(os.environ.get("PORT", "10000"))
MCP_SECRET = os.environ.get("MCP_SECRET", "")
YANDEX_API = "https://cloud-api.yandex.net/v1/disk"
APP_VERSION = "v6-diagnostic-return-20260830"

mcp = FastMCP(
    "Yandex Disk",
    instructions=(
        "Read/write access to the user's Yandex Disk. "
        "Use list_folder/search_files to browse; read_excel/read_text_file to inspect; "
        "use write tools only when the user asks to modify files. "
        "Destructive operations should keep backups when available; delete_file defaults to Trash."
    ),
    host="0.0.0.0",
    port=PORT,
    stateless_http=True,
    json_response=True,
    streamable_http_path=f"/mcp/{MCP_SECRET}" if MCP_SECRET else "/mcp",
)

def yandex_headers() -> dict[str, str]:
    return {"Authorization": f"OAuth {YANDEX_TOKEN}"}

def _raise(r: requests.Response) -> None:
    if not r.ok:
        raise RuntimeError(f"Yandex Disk API error {r.status_code}: {r.text[:1000]}")

def _normalize_path(path: str) -> str:
    if not path:
        return "/"
    if path.startswith("disk:"):
        path = path[5:]
    if not path.startswith("/"):
        path = "/" + path
    return posixpath.normpath(path)

def _parent(path: str) -> str:
    p = _normalize_path(path)
    parent = posixpath.dirname(p)
    return parent if parent else "/"

def _basename(path: str) -> str:
    return posixpath.basename(_normalize_path(path))

def _join(folder: str, name: str) -> str:
    return _normalize_path(posixpath.join(_normalize_path(folder), name))

def _resource(path: str, fields: str | None = None) -> dict[str, Any]:
    params: dict[str, Any] = {"path": _normalize_path(path)}
    if fields:
        params["fields"] = fields
    r = requests.get(f"{YANDEX_API}/resources", headers=yandex_headers(), params=params, timeout=30)
    _raise(r)
    return r.json()

def _resource_exists(path: str) -> bool:
    r = requests.get(
        f"{YANDEX_API}/resources",
        headers=yandex_headers(),
        params={"path": _normalize_path(path), "fields": "type,name,path"},
        timeout=30,
    )
    if r.status_code == 404:
        return False
    _raise(r)
    return True

def _download_bytes(path: str) -> bytes:
    r = requests.get(
        f"{YANDEX_API}/resources/download",
        headers=yandex_headers(),
        params={"path": _normalize_path(path)},
        timeout=30,
    )
    _raise(r)
    href = r.json()["href"]
    d = requests.get(href, timeout=120)
    _raise(d)
    return d.content

def _upload_bytes(path: str, content: bytes, overwrite: bool = False) -> dict[str, Any]:
    path = _normalize_path(path)

    # 1) Ask Yandex Disk for a temporary upload URL.
    r = requests.get(
        f"{YANDEX_API}/resources/upload",
        headers=yandex_headers(),
        params={"path": path, "overwrite": overwrite},
        timeout=30,
    )
    _raise(r)

    upload_info = r.json()
    href = upload_info.get("href")
    method = str(upload_info.get("method") or "PUT").upper()

    if not href:
        raise RuntimeError(f"Yandex Disk did not return an upload URL for {path}")

    # 2) Upload the raw bytes to the temporary URL.
    upload_headers = {
        "Content-Type": "application/octet-stream",
        "Content-Length": str(len(content)),
    }

    if method == "PUT":
        u = requests.put(
            href,
            data=io.BytesIO(content),
            headers=upload_headers,
            timeout=180,
            allow_redirects=True,
        )
    elif method == "POST":
        u = requests.post(
            href,
            data=io.BytesIO(content),
            headers=upload_headers,
            timeout=180,
            allow_redirects=True,
        )
    else:
        raise RuntimeError(f"Unsupported upload method from Yandex Disk: {method}")

    _raise(u)

    # 3) Verify that the file actually appeared on Disk.
    # Small files normally appear immediately, but we retry briefly.
    verified = False
    last_error = None
    for _ in range(10):
        try:
            if _resource_exists(path):
                verified = True
                break
        except Exception as exc:
            last_error = str(exc)
        time.sleep(0.5)

    if not verified:
        raise RuntimeError(
            f"Upload request succeeded, but file was not found on Yandex Disk: {path}. "
            f"Upload status={u.status_code}, response={u.text[:500]!r}, "
            f"last_check_error={last_error!r}"
        )

    return {
        "ok": True,
        "path": path,
        "bytes_uploaded": len(content),
        "overwrite": overwrite,
        "upload_status": u.status_code,
        "verified": True,
    }

def _wait_operation(href: str | None, timeout_seconds: int = 60) -> dict[str, Any]:
    if not href:
        return {"status": "success"}
    deadline = time.time() + timeout_seconds
    while time.time() < deadline:
        r = requests.get(href, headers=yandex_headers(), timeout=30)
        _raise(r)
        data = r.json()
        if data.get("status") in ("success", "failed"):
            return data
        time.sleep(1)
    return {"status": "in-progress", "href": href}

def _operation_result(r: requests.Response) -> dict[str, Any]:
    _raise(r)
    if r.status_code == 202:
        return _wait_operation(r.json().get("href"))
    if not r.content:
        return {"status": "success"}
    try:
        return r.json()
    except Exception:
        return {"status": "success"}

def _timestamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")

def _backup_path(path: str) -> str:
    path = _normalize_path(path)
    folder = _parent(path)
    name = _basename(path)
    root, ext = posixpath.splitext(name)
    return _join(folder, f"{root}_backup_{_timestamp()}{ext}")

def _copy_internal(source_path: str, destination_path: str, overwrite: bool = False) -> dict[str, Any]:
    r = requests.post(
        f"{YANDEX_API}/resources/copy",
        headers=yandex_headers(),
        params={"from": _normalize_path(source_path), "path": _normalize_path(destination_path), "overwrite": str(overwrite).lower()},
        timeout=30,
    )
    result = _operation_result(r)
    return {"ok": result.get("status") != "failed", "source": _normalize_path(source_path), "destination": _normalize_path(destination_path), "operation": result}

def _backup_if_requested(path: str, make_backup: bool) -> str | None:
    path = _normalize_path(path)
    if not make_backup or not _resource_exists(path):
        return None
    backup = _backup_path(path)
    _copy_internal(path, backup, overwrite=False)
    return backup

def _save_workbook(path: str, workbook, make_backup: bool = True) -> dict[str, Any]:
    path = _normalize_path(path)
    backup = _backup_if_requested(path, make_backup)
    out = io.BytesIO()
    workbook.save(out)
    result = _upload_bytes(path, out.getvalue(), overwrite=True)
    result["backup_path"] = backup
    return result

def _load_workbook_from_disk(path: str):
    return load_workbook(io.BytesIO(_download_bytes(path)))

def _list_folder_internal(path: str = "/") -> list[dict[str, Any]]:
    path = _normalize_path(path)
    items: list[dict[str, Any]] = []
    offset = 0
    limit = 1000
    while True:
        r = requests.get(f"{YANDEX_API}/resources", headers=yandex_headers(), params={"path": path, "limit": limit, "offset": offset}, timeout=30)
        _raise(r)
        embedded = r.json().get("_embedded", {})
        batch = embedded.get("items", [])
        for i in batch:
            items.append({"name": i.get("name"), "type": i.get("type"), "path": i.get("path"), "size": i.get("size"), "mime_type": i.get("mime_type"), "modified": i.get("modified"), "created": i.get("created")})
        total = embedded.get("total", len(items))
        offset += len(batch)
        if not batch or offset >= total:
            break
    return items


@mcp.tool()
def list_folder(path: str = "/") -> list[dict[str, Any]]:
    return _list_folder_internal(path)

@mcp.tool()
def get_file_info(path: str) -> dict[str, Any]:
    return _resource(path, fields="name,type,path,size,mime_type,created,modified,md5,sha256,revision")

@mcp.tool()
def file_exists(path: str) -> bool:
    try:
        return _resource(path, fields="type").get("type") == "file"
    except Exception:
        return False

@mcp.tool()
def folder_exists(path: str) -> bool:
    try:
        return _resource(path, fields="type").get("type") == "dir"
    except Exception:
        return False

@mcp.tool()
def search_files(query: str, start_path: str = "/", max_results: int = 100, max_depth: int = 8) -> list[dict[str, Any]]:
    query_norm = query.casefold().strip()
    results: list[dict[str, Any]] = []
    queue: list[tuple[str, int]] = [(_normalize_path(start_path), 0)]
    while queue and len(results) < max_results:
        folder, depth = queue.pop(0)
        try:
            items = _list_folder_internal(folder)
        except Exception:
            continue
        for item in items:
            name = str(item.get("name") or "")
            if query_norm in name.casefold():
                results.append(item)
                if len(results) >= max_results:
                    break
            if item.get("type") == "dir" and depth < max_depth:
                raw_path = item.get("path") or ""
                if isinstance(raw_path, str) and raw_path.startswith("disk:"):
                    raw_path = raw_path[5:]
                queue.append((_normalize_path(raw_path), depth + 1))
    return results

@mcp.tool()
def get_file_download_url(path: str) -> dict[str, str]:
    path = _normalize_path(path)
    r = requests.get(f"{YANDEX_API}/resources/download", headers=yandex_headers(), params={"path": path}, timeout=30)
    _raise(r)
    return {"path": path, "download_url": r.json()["href"]}

@mcp.tool()
def read_text_file(path: str, max_chars: int = 50000, encoding: str = "utf-8") -> dict[str, Any]:
    text = _download_bytes(path).decode(encoding, errors="replace")
    return {"path": _normalize_path(path), "chars_total": len(text), "truncated": len(text) > max_chars, "text": text[:max_chars]}

@mcp.tool()
def read_excel(path: str, sheet: str = "", max_rows: int = 200) -> dict[str, Any]:
    wb = _load_workbook_from_disk(path)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = []
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if idx > max_rows:
            break
        rows.append(list(row))
    return {"path": _normalize_path(path), "sheets": wb.sheetnames, "selected_sheet": ws.title, "rows_returned": len(rows), "rows": rows}

@mcp.tool()
def read_excel_sheet_names(path: str) -> list[str]:
    return _load_workbook_from_disk(path).sheetnames

@mcp.tool()
def create_folder(path: str) -> dict[str, Any]:
    path = _normalize_path(path)
    r = requests.put(f"{YANDEX_API}/resources", headers=yandex_headers(), params={"path": path}, timeout=30)
    if r.status_code == 409 and _resource_exists(path):
        return {"ok": True, "path": path, "already_exists": True}
    _raise(r)
    return {"ok": True, "path": path, "already_exists": False}

@mcp.tool()
def copy_file(source_path: str, destination_path: str, overwrite: bool = False) -> dict[str, Any]:
    return _copy_internal(source_path, destination_path, overwrite)

def _move_diagnostic(source_path: str, destination_path: str, overwrite: bool = False, execute: bool = True) -> dict[str, Any]:
    """Run or inspect a Yandex Disk move without hiding any diagnostic detail.

    This helper deliberately avoids raising for Yandex API errors so the MCP client
    receives the HTTP status/body directly.
    """
    result: dict[str, Any] = {
        "ok": False,
        "bridge_version": APP_VERSION,
        "execute": execute,
        "source": None,
        "destination": None,
        "overwrite": bool(overwrite),
        "stages": [],
    }
    try:
        source_path = _normalize_path(source_path)
        destination_path = _normalize_path(destination_path)
        result["source"] = source_path
        result["destination"] = destination_path
        result["stages"].append({"stage": "normalized"})

        source_exists_before = _resource_exists(source_path)
        destination_exists_before = _resource_exists(destination_path)
        result["source_exists_before"] = source_exists_before
        result["destination_exists_before"] = destination_exists_before
        result["stages"].append({
            "stage": "precheck",
            "source_exists": source_exists_before,
            "destination_exists": destination_exists_before,
        })

        if not source_exists_before:
            result["error_type"] = "FileNotFoundError"
            result["error"] = f"Source does not exist: {source_path}"
            return result

        if source_path == destination_path:
            result.update({"ok": True, "already_at_destination": True, "verified": True})
            result["stages"].append({"stage": "no_op_same_path"})
            return result

        params = {
            "from": source_path,
            "path": destination_path,
            "overwrite": str(bool(overwrite)).lower(),
        }
        result["request"] = {
            "method": "POST",
            "endpoint": "/resources/move",
            "params": params,
        }

        if not execute:
            result["ok"] = True
            result["stages"].append({"stage": "dry_run_ready"})
            return result

        r = requests.post(
            f"{YANDEX_API}/resources/move",
            headers=yandex_headers(),
            params=params,
            timeout=30,
        )
        body = (r.text or "")[:4000]
        result["yandex_http_status"] = r.status_code
        result["yandex_response_body"] = body
        result["stages"].append({"stage": "move_http_response", "status": r.status_code})

        parsed: Any = None
        if r.content:
            try:
                parsed = r.json()
            except Exception:
                parsed = None
        if parsed is not None:
            result["yandex_response_json"] = parsed

        if not r.ok:
            result["error_type"] = "YandexDiskHTTPError"
            result["error"] = f"Yandex Disk API returned HTTP {r.status_code}"
            return result

        operation: dict[str, Any] = {"status": "success"}
        if r.status_code == 202:
            href = parsed.get("href") if isinstance(parsed, dict) else None
            result["operation_href_present"] = bool(href)
            if href:
                deadline = time.time() + 30
                poll_count = 0
                while time.time() < deadline:
                    poll_count += 1
                    pr = requests.get(href, headers=yandex_headers(), timeout=30)
                    pbody = (pr.text or "")[:4000]
                    try:
                        pdata = pr.json() if pr.content else {}
                    except Exception:
                        pdata = {"raw_body": pbody}
                    operation = {
                        "http_status": pr.status_code,
                        "body": pdata,
                        "poll_count": poll_count,
                    }
                    status = pdata.get("status") if isinstance(pdata, dict) else None
                    if not pr.ok or status in ("success", "failed"):
                        break
                    time.sleep(0.5)
            else:
                operation = {"status": "unknown", "reason": "202 response without href"}
        elif isinstance(parsed, dict) and parsed:
            operation = parsed

        result["operation"] = operation
        result["stages"].append({"stage": "operation_checked"})

        destination_exists_after = _resource_exists(destination_path)
        source_exists_after = _resource_exists(source_path)
        result["source_exists_after"] = source_exists_after
        result["destination_exists_after"] = destination_exists_after
        result["verified"] = destination_exists_after and not source_exists_after
        result["stages"].append({
            "stage": "postcheck",
            "source_exists": source_exists_after,
            "destination_exists": destination_exists_after,
        })

        if result["verified"]:
            result["ok"] = True
        else:
            result["error_type"] = "MoveVerificationError"
            result["error"] = "Yandex request completed but the expected final file state was not observed."
        return result
    except Exception as exc:
        result["error_type"] = type(exc).__name__
        result["error"] = str(exc)
        result["traceback_tail"] = traceback.format_exc().splitlines()[-12:]
        return result


def _move_internal(source_path: str, destination_path: str, overwrite: bool = False) -> dict[str, Any]:
    result = _move_diagnostic(source_path, destination_path, overwrite, execute=True)
    if not result.get("ok"):
        raise RuntimeError(f"Move failed: {result}")
    return result


@mcp.tool()
def bridge_version() -> dict[str, Any]:
    return {"ok": True, "bridge_version": APP_VERSION}


@mcp.tool()
def diagnose_move(source_path: str, destination_path: str, overwrite: bool = False, execute: bool = False) -> dict[str, Any]:
    """Return move diagnostics directly to ChatGPT. Set execute=true to perform the move."""
    return _move_diagnostic(source_path, destination_path, overwrite, execute=execute)


@mcp.tool()
def move_file(source_path: str, destination_path: str, overwrite: bool = False) -> dict[str, Any]:
    return _move_diagnostic(source_path, destination_path, overwrite, execute=True)


@mcp.tool()
def rename_file(path: str, new_name: str, overwrite: bool = False) -> dict[str, Any]:
    result: dict[str, Any] = {
        "ok": False,
        "bridge_version": APP_VERSION,
        "path": path,
        "new_name": new_name,
        "overwrite": bool(overwrite),
    }
    try:
        normalized_path = _normalize_path(path)
        clean_name = new_name.strip()
        result["path"] = normalized_path
        result["new_name"] = clean_name
        if not clean_name or "/" in clean_name or "\\" in clean_name:
            result["error_type"] = "ValueError"
            result["error"] = "new_name must be a basename only"
            return result

        destination_path = _join(_parent(normalized_path), clean_name)
        result["destination"] = destination_path
        move_result = _move_diagnostic(normalized_path, destination_path, overwrite, execute=True)
        result["move"] = move_result
        result["ok"] = bool(move_result.get("ok"))
        if not result["ok"]:
            result["error_type"] = move_result.get("error_type", "RenameFailed")
            result["error"] = move_result.get("error", "Rename failed")
        return result
    except Exception as exc:
        result["error_type"] = type(exc).__name__
        result["error"] = str(exc)
        result["traceback_tail"] = traceback.format_exc().splitlines()[-12:]
        return result


@mcp.tool()
def backup_file(path: str) -> dict[str, Any]:
    destination = _backup_path(path)
    result = _copy_internal(path, destination, False)
    result["backup_path"] = destination
    return result

@mcp.tool()
def delete_file(path: str, permanently: bool = False) -> dict[str, Any]:
    path = _normalize_path(path)
    r = requests.delete(
        f"{YANDEX_API}/resources",
        headers=yandex_headers(),
        params={"path": path, "permanently": str(permanently).lower()},
        timeout=30,
    )
    op = _operation_result(r)
    return {"ok": op.get("status") != "failed", "path": path, "permanently": permanently, "operation": op}

@mcp.tool()
def list_trash(limit: int = 100) -> list[dict[str, Any]]:
    r = requests.get(f"{YANDEX_API}/trash/resources", headers=yandex_headers(), params={"limit": min(max(limit, 1), 1000)}, timeout=30)
    _raise(r)
    items = r.json().get("_embedded", {}).get("items", [])
    return [{"name": i.get("name"), "type": i.get("type"), "path": i.get("path"), "deleted": i.get("deleted"), "origin_path": i.get("origin_path")} for i in items]

@mcp.tool()
def restore_from_trash(trash_path: str, new_name: str = "", overwrite: bool = False) -> dict[str, Any]:
    params: dict[str, Any] = {"path": trash_path, "overwrite": str(overwrite).lower()}
    if new_name:
        params["name"] = new_name
    r = requests.put(f"{YANDEX_API}/trash/resources/restore", headers=yandex_headers(), params=params, timeout=30)
    op = _operation_result(r)
    return {"ok": op.get("status") != "failed", "trash_path": trash_path, "new_name": new_name or None, "overwrite": overwrite, "operation": op}

@mcp.tool()
def upload_file(path: str, content_base64: str, overwrite: bool = False) -> dict[str, Any]:
    try:
        data = base64.b64decode(content_base64, validate=True)
    except Exception as exc:
        raise ValueError(f"Invalid base64 content: {exc}") from exc
    return _upload_bytes(path, data, overwrite)

@mcp.tool()
def overwrite_file(path: str, content_base64: str, make_backup: bool = True) -> dict[str, Any]:
    path = _normalize_path(path)
    backup = _backup_if_requested(path, make_backup)
    try:
        data = base64.b64decode(content_base64, validate=True)
    except Exception as exc:
        raise ValueError(f"Invalid base64 content: {exc}") from exc
    result = _upload_bytes(path, data, True)
    result["backup_path"] = backup
    return result

@mcp.tool()
def write_text_file(path: str, text: str, overwrite: bool = False, encoding: str = "utf-8", make_backup: bool = True) -> dict[str, Any]:
    path = _normalize_path(path)
    backup = _backup_if_requested(path, make_backup) if overwrite else None
    result = _upload_bytes(path, text.encode(encoding), overwrite)
    result["backup_path"] = backup
    return result

@mcp.tool()
def append_text_file(path: str, text: str, encoding: str = "utf-8", make_backup: bool = True) -> dict[str, Any]:
    path = _normalize_path(path)
    existing = _download_bytes(path).decode(encoding, errors="replace")
    backup = _backup_if_requested(path, make_backup)
    result = _upload_bytes(path, (existing + text).encode(encoding), True)
    result["backup_path"] = backup
    return result

@mcp.tool()
def replace_text(path: str, old_text: str, new_text: str, replace_all: bool = True, encoding: str = "utf-8", make_backup: bool = True) -> dict[str, Any]:
    path = _normalize_path(path)
    existing = _download_bytes(path).decode(encoding, errors="replace")
    count = existing.count(old_text)
    if count == 0:
        return {"ok": False, "path": path, "replacements": 0, "message": "old_text not found"}
    updated = existing.replace(old_text, new_text) if replace_all else existing.replace(old_text, new_text, 1)
    backup = _backup_if_requested(path, make_backup)
    result = _upload_bytes(path, updated.encode(encoding), True)
    result.update({"backup_path": backup, "replacements": count if replace_all else 1})
    return result

@mcp.tool()
def update_excel_cell(path: str, sheet: str, cell: str, value: Any, make_backup: bool = True) -> dict[str, Any]:
    wb = _load_workbook_from_disk(path)
    if sheet not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet}")
    ws = wb[sheet]
    old_value = ws[cell].value
    ws[cell] = value
    result = _save_workbook(path, wb, make_backup)
    result.update({"sheet": sheet, "cell": cell, "old_value": old_value, "new_value": value})
    return result

def _update_excel_range_internal(path: str, sheet: str, start_cell: str, values: list[list[Any]], make_backup: bool = True) -> dict[str, Any]:
    wb = _load_workbook_from_disk(path)
    if sheet not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet}")
    ws = wb[sheet]
    anchor = ws[start_cell]
    for r_idx, row_values in enumerate(values):
        for c_idx, value in enumerate(row_values):
            ws.cell(row=anchor.row + r_idx, column=anchor.column + c_idx, value=value)
    result = _save_workbook(path, wb, make_backup)
    result.update({"sheet": sheet, "start_cell": start_cell, "rows_written": len(values), "max_columns_written": max((len(r) for r in values), default=0)})
    return result


@mcp.tool()
def update_excel_range(path: str, sheet: str, start_cell: str, values: list[list[Any]], make_backup: bool = True) -> dict[str, Any]:
    return _update_excel_range_internal(path, sheet, start_cell, values, make_backup)

@mcp.tool()
def append_excel_row(path: str, sheet: str, values: list[Any], make_backup: bool = True) -> dict[str, Any]:
    wb = _load_workbook_from_disk(path)
    if sheet not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet}")
    ws = wb[sheet]
    ws.append(values)
    row_number = ws.max_row
    result = _save_workbook(path, wb, make_backup)
    result.update({"sheet": sheet, "row_number": row_number, "values": values})
    return result

@mcp.tool()
def write_excel_table(path: str, sheet: str, start_cell: str, rows: list[list[Any]], make_backup: bool = True) -> dict[str, Any]:
    return _update_excel_range_internal(path, sheet, start_cell, rows, make_backup)

@mcp.tool()
def add_excel_sheet(path: str, sheet_name: str, index: int = -1, make_backup: bool = True) -> dict[str, Any]:
    wb = _load_workbook_from_disk(path)
    if sheet_name in wb.sheetnames:
        raise ValueError(f"Sheet already exists: {sheet_name}")
    if index < 0:
        wb.create_sheet(title=sheet_name)
    else:
        wb.create_sheet(title=sheet_name, index=index)
    result = _save_workbook(path, wb, make_backup)
    result.update({"sheet_added": sheet_name, "sheet_names": wb.sheetnames})
    return result

@mcp.tool()
def rename_excel_sheet(path: str, old_name: str, new_name: str, make_backup: bool = True) -> dict[str, Any]:
    wb = _load_workbook_from_disk(path)
    if old_name not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {old_name}")
    if new_name in wb.sheetnames:
        raise ValueError(f"Sheet already exists: {new_name}")
    wb[old_name].title = new_name
    result = _save_workbook(path, wb, make_backup)
    result.update({"old_name": old_name, "new_name": new_name, "sheet_names": wb.sheetnames})
    return result

@mcp.tool()
def delete_excel_sheet(path: str, sheet_name: str, make_backup: bool = True) -> dict[str, Any]:
    wb = _load_workbook_from_disk(path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}")
    if len(wb.sheetnames) <= 1:
        raise ValueError("Cannot delete the only worksheet in the workbook")
    del wb[sheet_name]
    result = _save_workbook(path, wb, make_backup)
    result.update({"sheet_deleted": sheet_name, "sheet_names": wb.sheetnames})
    return result

if __name__ == "__main__":
    mcp.run(transport="streamable-http")
